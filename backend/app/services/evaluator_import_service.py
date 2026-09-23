"""
evaluator_import_service.py — Importación de Servicio_WOW_2026_Evaluadores.xlsx.

Hojas que se leen:
    Cronograma      → fechas, áreas que evalúan y muestras, por 'No.'
    Evaluadores     → una fila por persona asignada a una lista
                      (upsert de ScheduleEntry, Employee, EvaluatorArea, EvaluationAssignment)
    Resumen_Muestra → SOLO la columna 'Muestra requerida' (celda editable, dato
                      de entrada: en Caja la muestra es por cajera y no sale del
                      Cronograma). El resto de esa hoja es derivado y no se importa.

No se importan: Matriz_Participacion, Criterios_Seleccion, Formularios,
Resultados Generales ni las hojas de detalle por depto (GH, Compras, ...).

Unidad = LISTA ('Departamento evaluado', ej: 'CAJA – Gurabo'). Las fechas se
toman del Cronograma por 'No.' — Caja es una sola fila del Cronograma que
aplica a sus 5 listas.

Es un reemplazo manual: se corre después de cada ejecución del script offline.
Las asignaciones PENDIENTES de una lista que ya no aparecen en el Excel se
eliminan (el script las re-sorteó); las COMPLETAS se conservan.
"""

import io
import logging
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from difflib import SequenceMatcher
from typing import Optional

import openpyxl
from sqlalchemy.orm import Session

from app.models.evaluator_models import (
    AssignmentRole, AssignmentStatus, Employee, EvaluationAssignment,
    EvaluatorArea, ScheduleEntry,
)
from app.models.survey_wow_models import SurveyDepartment, SurveyForm

logger = logging.getLogger(__name__)


# ─────────────────────────────────────────────────────────────────────────────
# NORMALIZACIÓN DE NOMBRES (compartida con survey_wow_import_service)
# ─────────────────────────────────────────────────────────────────────────────

def norm_name(txt) -> str:
    """Igual que norm() de asignar_evaluadores_wow.py: sin tildes, solo letras, MAYÚSCULAS."""
    if txt is None:
        return ""
    txt = unicodedata.normalize("NFKD", str(txt)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", re.sub(r"[^A-Za-z ]", " ", txt)).strip().upper()


def norm_label(txt) -> str:
    """Normaliza encabezados / nombres de lista (conserva dígitos, colapsa guiones y espacios)."""
    if txt is None:
        return ""
    txt = unicodedata.normalize("NFKD", str(txt).replace("\xa0", " ")).encode("ascii", "ignore").decode()
    txt = re.sub(r"[–—-]", " - ", txt)
    return re.sub(r"\s+", " ", txt).strip().lower()


def name_matches(candidate: str, target_key: str) -> bool:
    """
    Emparejamiento tolerante, misma regla que emparejar() del script:
    cada token (>2 letras) del candidato se compara con los tokens del
    nombre del listado (SequenceMatcher ≥ 0.8); el primer token debe
    calzar y al menos min(2, n) tokens en total.
    """
    tok = [t for t in norm_name(candidate).split() if len(t) > 2]
    if not tok:
        return False
    k = target_key.split()
    ok = [any(SequenceMatcher(None, t, x).ratio() >= 0.8 for x in k) for t in tok]
    return ok[0] and sum(ok) >= min(2, len(tok))


# ─────────────────────────────────────────────────────────────────────────────
# PARSERS DE CELDAS
# ─────────────────────────────────────────────────────────────────────────────

_MESES = {"ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6,
          "jul": 7, "ago": 8, "sep": 9, "sept": 9, "oct": 10, "nov": 11, "dic": 12}


def _to_date(val) -> Optional[date]:
    """Acepta datetime/date de Excel o texto tipo '05-ene-2027' / '2026-09-28'."""
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    s = str(val).strip().lower()
    m = re.match(r"^(\d{1,2})[-/ ]([a-z]+)\.?[-/ ](\d{4})$", s)
    if m:
        mes = _MESES.get(m.group(2)) or _MESES.get(m.group(2)[:3])
        if mes:
            return date(int(m.group(3)), mes, int(m.group(1)))
    try:
        return date.fromisoformat(s[:10])
    except ValueError:
        return None


def _leading_int(val) -> Optional[int]:
    """'35 colaboradores' → 35 ; 'N/A' / None → None ; 12 → 12."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return int(val)
    m = re.match(r"^\s*(\d+)", str(val))
    return int(m.group(1)) if m else None


def _str(val) -> Optional[str]:
    if val is None:
        return None
    s = re.sub(r"\s+", " ", str(val).replace("\xa0", " ")).strip()
    return s or None


def _header_index(headers: list, *aliases: str) -> Optional[int]:
    """Índice de la primera columna cuyo encabezado normalizado coincide con algún alias."""
    normed = [norm_label(h) for h in headers]
    for alias in aliases:
        a = norm_label(alias)
        for i, h in enumerate(normed):
            if h == a:
                return i
    for alias in aliases:
        a = norm_label(alias)
        for i, h in enumerate(normed):
            if h.startswith(a):
                return i
    return None


# ─────────────────────────────────────────────────────────────────────────────
# RESULTADO
# ─────────────────────────────────────────────────────────────────────────────

@dataclass
class EvaluatorImportResult:
    listas_creadas:          int = 0
    listas_actualizadas:     int = 0
    colaboradores_creados:   int = 0
    asignaciones_creadas:    int = 0
    asignaciones_actualizadas: int = 0
    asignaciones_eliminadas: int = 0
    departamentos_creados:   list[str]  = field(default_factory=list)
    advertencias:            list[str]  = field(default_factory=list)
    errores:                 list[dict] = field(default_factory=list)


# ─────────────────────────────────────────────────────────────────────────────
# IMPORTACIÓN
# ─────────────────────────────────────────────────────────────────────────────

def _read_cronograma(wb) -> dict[int, dict]:
    if "Cronograma" not in wb.sheetnames:
        raise ValueError("El Excel no tiene la hoja 'Cronograma'.")
    rows = list(wb["Cronograma"].iter_rows(values_only=True))
    headers = list(rows[0])
    idx = {
        "no":       _header_index(headers, "No."),
        "send":     _header_index(headers, "Envío de encuesta"),
        "tab":      _header_index(headers, "Tabulación"),
        "disc":     _header_index(headers, "Divulgación"),
        "report":   _header_index(headers, "Entrega informe general"),
        "areas":    _header_index(headers, "Áreas que evalúan"),
        "int":      _header_index(headers, "Muestra interna"),
        "ext":      _header_index(headers, "Muestra Externa"),
    }
    if idx["no"] is None:
        raise ValueError("La hoja 'Cronograma' no tiene la columna 'No.'.")

    def cell(r, key):
        i = idx[key]
        return r[i] if i is not None and i < len(r) else None

    out: dict[int, dict] = {}
    for r in rows[1:]:
        no = _leading_int(cell(r, "no"))
        if no is None:
            continue
        out[no] = {
            "send_date":            _to_date(cell(r, "send")),
            "tabulation_date":      _to_date(cell(r, "tab")),
            "disclosure_date":      _to_date(cell(r, "disc")),
            "report_delivery_date": _to_date(cell(r, "report")),
            "evaluator_areas_raw":  _str(cell(r, "areas")),
            "internal_sample_raw":  _str(cell(r, "int")),
            "external_sample_raw":  _str(cell(r, "ext")),
            "internal_sample_size": _leading_int(cell(r, "int")),
            "external_sample_size": _leading_int(cell(r, "ext")),
        }
    return out


def _read_muestra_requerida(wb) -> dict[str, int]:
    """'Muestra requerida' por lista, desde Resumen_Muestra (si existe)."""
    if "Resumen_Muestra" not in wb.sheetnames:
        return {}
    rows = list(wb["Resumen_Muestra"].iter_rows(values_only=True))
    if not rows:
        return {}
    headers = list(rows[0])
    i_lista = _header_index(headers, "Departamento evaluado")
    i_req = _header_index(headers, "Muestra requerida")
    if i_lista is None or i_req is None:
        return {}
    out = {}
    for r in rows[1:]:
        lista, req = _str(r[i_lista]), _leading_int(r[i_req])
        if lista and req is not None and norm_label(lista) != "total":
            out[norm_label(lista)] = req
    return out


def importar_evaluadores_desde_excel(
    file_bytes: bytes,
    db: Session,
    cycle_id: int,
) -> EvaluatorImportResult:
    result = EvaluatorImportResult()

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as e:
        result.errores.append({"fila": "N/A", "error": f"No se pudo leer el Excel: {e}"})
        return result

    try:
        cronograma = _read_cronograma(wb)
    except ValueError as e:
        result.errores.append({"fila": "N/A", "error": str(e)})
        return result
    muestra_req = _read_muestra_requerida(wb)

    if "Evaluadores" not in wb.sheetnames:
        result.errores.append({"fila": "N/A", "error": "El Excel no tiene la hoja 'Evaluadores'."})
        return result
    rows = list(wb["Evaluadores"].iter_rows(values_only=True))
    headers = list(rows[0])
    col = {
        "no":        _header_index(headers, "No."),
        "lista":     _header_index(headers, "Departamento evaluado"),
        "envio":     _header_index(headers, "Envío"),
        "area_eval": _header_index(headers, "Área evaluadora (cronograma)", "Área evaluadora"),
        "tipo":      _header_index(headers, "Tipo"),
        "nombre":    _header_index(headers, "Colaborador"),
        "puesto":    _header_index(headers, "Puesto"),
        "area":      _header_index(headers, "Área"),
        "subdepto":  _header_index(headers, "Sub-departamento"),
        "ubic":      _header_index(headers, "Ubicación"),
        "respondio": _header_index(headers, "Respondió"),
    }
    faltantes = [k for k in ("lista", "tipo", "nombre") if col[k] is None]
    if faltantes:
        result.errores.append({
            "fila": "N/A",
            "error": f"Hoja 'Evaluadores' sin columnas requeridas: {faltantes}. Encabezados: {headers}",
        })
        return result

    def cell(r, key):
        i = col[key]
        return r[i] if i is not None and i < len(r) else None

    # ── Catálogos en memoria ─────────────────────────────────────────────────
    forms_by_list = {
        norm_label(f.list_name): f
        for f in db.query(SurveyForm).filter(SurveyForm.cycle_id == cycle_id, SurveyForm.list_name.isnot(None))
    }
    depts = {norm_label(d.name): d for d in db.query(SurveyDepartment).all()}
    areas = {a.name: a for a in db.query(EvaluatorArea).all()}
    employees = {e.name_key: e for e in db.query(Employee).all()}
    entries = {
        norm_label(e.list_name): e
        for e in db.query(ScheduleEntry).filter(ScheduleEntry.cycle_id == cycle_id)
    }

    def get_area(name: Optional[str]) -> Optional[EvaluatorArea]:
        if not name:
            return None
        if name not in areas:
            areas[name] = EvaluatorArea(name=name)
            db.add(areas[name])
        return areas[name]

    def resolve_dept_branch(list_name: str) -> tuple[SurveyDepartment, Optional[str]]:
        form = forms_by_list.get(norm_label(list_name))
        if form:
            return form.department, form.branch
        # Lista sin formulario en el catálogo: 'DEPTO – Sucursal' → crear depto si no existe
        partes = re.split(r"\s+[–—-]\s+", list_name, maxsplit=1)
        dept_name = partes[0].strip().title()
        branch = partes[1].strip() if len(partes) > 1 else None
        dept = depts.get(norm_label(dept_name))
        if not dept:
            dept = SurveyDepartment(name=dept_name, is_active=True)
            db.add(dept)
            db.flush()
            depts[norm_label(dept_name)] = dept
            result.departamentos_creados.append(dept_name)
        result.advertencias.append(
            f"La lista '{list_name}' no tiene formulario interno en el catálogo; "
            f"se asignó al departamento '{dept.name}'."
        )
        return dept, branch

    # ── Filas de Evaluadores ─────────────────────────────────────────────────
    seen_per_entry: dict[int, set[int]] = {}
    touched_entries: dict[str, ScheduleEntry] = {}

    for n_fila, r in enumerate(rows[1:], start=2):
        list_name = _str(cell(r, "lista"))
        nombre = _str(cell(r, "nombre"))
        if not list_name or not nombre:
            continue
        try:
            key_list = norm_label(list_name)
            entry = touched_entries.get(key_list)
            if entry is None:
                no = _leading_int(cell(r, "no"))
                crono = cronograma.get(no, {}) if no is not None else {}
                if no is not None and not crono:
                    result.advertencias.append(f"Lista '{list_name}': No.={no} no está en el Cronograma.")
                dept, branch = resolve_dept_branch(list_name)
                entry = entries.get(key_list)
                if entry is None:
                    entry = ScheduleEntry(cycle_id=cycle_id, list_name=list_name)
                    db.add(entry)
                    entries[key_list] = entry
                    result.listas_creadas += 1
                else:
                    result.listas_actualizadas += 1
                entry.schedule_number = no
                entry.department_id = dept.id
                entry.branch = branch
                for k, v in crono.items():
                    setattr(entry, k, v)
                if entry.send_date is None:
                    entry.send_date = _to_date(cell(r, "envio"))
                if key_list in muestra_req:
                    entry.internal_sample_size = muestra_req[key_list]
                db.flush()
                touched_entries[key_list] = entry
                seen_per_entry[entry.id] = set()

            # Colaborador
            key = norm_name(nombre)
            emp = employees.get(key)
            if emp is None:
                emp = Employee(full_name=nombre, name_key=key, is_active=True)
                db.add(emp)
                employees[key] = emp
                result.colaboradores_creados += 1
            emp.position = _str(cell(r, "puesto")) or emp.position
            emp.sub_department = _str(cell(r, "subdepto")) or emp.sub_department
            emp.location = _str(cell(r, "ubic")) or emp.location
            area = get_area(_str(cell(r, "area")))
            if area is not None:
                emp.area = area
            db.flush()

            # Asignación
            tipo = norm_label(cell(r, "tipo"))
            role = AssignmentRole.SUPLENTE.value if tipo.startswith("suplente") else AssignmentRole.TITULAR.value
            respondio = norm_label(cell(r, "respondio"))
            status_excel = (
                AssignmentStatus.PENDIENTE.value
                if respondio in ("", "pendiente", "no", "false")
                else AssignmentStatus.COMPLETO.value
            )
            eval_area = get_area(_str(cell(r, "area_eval")))

            asg = db.query(EvaluationAssignment).filter(
                EvaluationAssignment.schedule_entry_id == entry.id,
                EvaluationAssignment.employee_id == emp.id,
            ).first()
            if asg is None:
                asg = EvaluationAssignment(
                    schedule_entry_id=entry.id, employee_id=emp.id,
                    role=role, status=status_excel, evaluator_area=eval_area,
                )
                db.add(asg)
                result.asignaciones_creadas += 1
            else:
                asg.role = role
                asg.evaluator_area = eval_area
                # Nunca "des-completar": si ya se marcó completo al importar respuestas, se respeta.
                if status_excel == AssignmentStatus.COMPLETO.value:
                    asg.status = status_excel
                result.asignaciones_actualizadas += 1
            db.flush()
            seen_per_entry[entry.id].add(emp.id)

        except Exception as e:
            db.rollback()
            result.errores.append({"fila": n_fila, "lista": list_name, "error": str(e)})
            logger.error(f"Error en fila {n_fila} de Evaluadores: {e}", exc_info=True)
            return result

    # ── Reemplazo: pendientes que el script ya no asigna ─────────────────────
    # Las listas cerradas no se tocan: pueden tener evaluadores agregados desde la app.
    for entry_id, emp_ids in seen_per_entry.items():
        if not emp_ids or db.get(ScheduleEntry, entry_id).is_closed:
            continue
        stale = db.query(EvaluationAssignment).filter(
            EvaluationAssignment.schedule_entry_id == entry_id,
            EvaluationAssignment.status == AssignmentStatus.PENDIENTE.value,
            EvaluationAssignment.employee_id.notin_(emp_ids),
        ).all()
        for asg in stale:
            db.delete(asg)
            result.asignaciones_eliminadas += 1

    try:
        db.commit()
    except Exception as e:
        db.rollback()
        result.errores.append({"fila": "commit", "error": str(e)})
        logger.error(f"Error en commit de evaluadores: {e}", exc_info=True)

    logger.info(
        f"Import evaluadores: {result.listas_creadas} listas nuevas, "
        f"{result.asignaciones_creadas} asignaciones nuevas, "
        f"{result.asignaciones_eliminadas} eliminadas, {len(result.errores)} errores"
    )
    return result
